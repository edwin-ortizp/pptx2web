"""Actividades interactivas: cita en notas → config → copia estable → player."""
import json
from pathlib import Path

import pytest
from pptx import Presentation

from pptx2web import config as config_mod
from pptx2web.metadata import extract, normalize_alias
from pptx2web.models import SlideAssets
from pptx2web.packager import build_manifest, cited_activities, package, update_activities
from pptx2web.validate import ValidationError


def _deck(tmp_path: Path, notes: list[str]) -> Path:
    prs = Presentation()
    for text in notes:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "T"
        if text:
            slide.notes_slide.notes_text_frame.text = text
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    return path


def _assets(n: int) -> list[SlideAssets]:
    return [SlideAssets(index=i, src=f"slides/s{i}.webp", thumb=f"thumbs/t{i}.webp")
            for i in range(1, n + 1)]


def test_normalize_alias():
    assert normalize_alias("Sprint Eventos") == "sprint-eventos"
    assert normalize_alias("  Ética_IA! ") == "etica-ia"


def test_activity_parsed_and_stripped_from_notes(tmp_path: Path):
    path = _deck(tmp_path, [
        "Nota previa\n[actividad: Sprint Eventos | Ordena los eventos]\nNota final",
        "[actividad: prompt-lab]",
        "sin actividad",
    ])
    deck = extract(path)
    a1 = deck.slides[0].activity
    assert a1.alias == "sprint-eventos" and a1.title == "Ordena los eventos"
    assert "actividad" not in deck.slides[0].notes_html
    assert "Nota previa" in deck.slides[0].notes_html
    assert deck.slides[1].activity.alias == "prompt-lab"
    assert deck.slides[1].activity.title is None
    assert deck.slides[1].notes_html is None
    assert deck.slides[2].activity is None


def test_activity_coexists_with_quiz(tmp_path: Path):
    path = _deck(tmp_path, ["[quiz]\n- a\n+ b\n[actividad: x]"])
    s = extract(path).slides[0]
    assert s.quiz is not None and s.activity.alias == "x"


def test_second_activity_on_slide_warns(tmp_path: Path):
    deck = extract(_deck(tmp_path, ["[actividad: a]\n[actividad: b]"]))
    assert deck.slides[0].activity.alias == "a"
    assert any("1 actividad" in w for w in deck.warnings)


def test_validate_activities():
    out = config_mod._validate_activities({"Mi Act": "x/mi.html",
                                           "web": {"file": "https://e.com/a"}})
    assert out["mi-act"]["file"] == "x/mi.html"
    assert out["web"]["file"] == "https://e.com/a"
    with pytest.raises(ValidationError):
        config_mod._validate_activities({"a": {"file": "a.pdf"}})
    with pytest.raises(ValidationError):
        config_mod._validate_activities(["a"])


def test_activity_status(tmp_path: Path):
    (tmp_path / "actividades").mkdir()
    (tmp_path / "actividades" / "auto.html").write_text("<p>auto</p>")
    (tmp_path / "reg.html").write_text("<p>reg</p>")
    acts = config_mod._validate_activities({"reg": "reg.html", "sobra": "nada.html"})
    rows = {r["alias"]: r for r in config_mod.activity_status(
        {"auto": [2], "reg": [3], "falta": [4]}, acts, tmp_path)}
    assert rows["auto"]["status"] == "ok" and rows["auto"]["auto"]
    assert rows["reg"]["status"] == "ok" and not rows["reg"]["auto"]
    assert rows["falta"]["status"] == "missing_file"
    assert rows["sobra"]["status"] == "unused"
    warnings = config_mod.activity_warnings(list(rows.values()))
    assert len(warnings) == 2


def _publish(tmp_path: Path, cfg: dict):
    path = _deck(tmp_path, ["[actividad: sprint]", "[actividad: falta]", "[actividad: web]"])
    deck = extract(path)
    manifest = build_manifest(deck, _assets(3), (1280, 720))
    player_config, _ = config_mod.resolve(cfg, None, 3)
    out = tmp_path / "out"
    out.mkdir()
    package(manifest, out, player_config, config_dir=tmp_path)
    return out, manifest


def test_package_copies_with_stable_name(tmp_path: Path):
    (tmp_path / "actividades").mkdir()
    src = tmp_path / "actividades" / "sprint.html"
    src.write_text("<p>v1</p>")
    out, manifest = _publish(tmp_path, {"activities": {"web": "https://e.com/a"}})

    assert manifest["slides"][0]["activity"] == {"alias": "sprint", "title": None}
    assert cited_activities(manifest) == {"sprint": [1], "falta": [2], "web": [3]}
    cfg = json.loads((out / "config.json").read_text(encoding="utf-8"))
    assert cfg["activities"]["sprint"]["src"] == "activities/sprint.html"
    assert cfg["activities"]["sprint"]["version"]
    assert cfg["activities"]["web"] == {"src": "https://e.com/a", "title": None, "version": None}
    assert "falta" not in cfg["activities"]
    assert (out / "activities" / "sprint.html").read_text() == "<p>v1</p>"
    assert "activities/sprint.html" in (out / "index.html").read_text(encoding="utf-8")


def test_update_activities_without_reconverting(tmp_path: Path):
    (tmp_path / "actividades").mkdir()
    src = tmp_path / "actividades" / "sprint.html"
    src.write_text("<p>v1</p>")
    out, _ = _publish(tmp_path, {})
    v1 = json.loads((out / "config.json").read_text())["activities"]["sprint"]["version"]
    (out / "activities" / "viejo.html").write_text("huérfano")

    src.write_text("<p>v2</p>")
    (tmp_path / "actividades" / "falta.html").write_text("<p>nueva</p>")
    update_activities(out, {}, tmp_path)

    cfg = json.loads((out / "config.json").read_text())
    assert cfg["activities"]["sprint"]["version"] != v1
    assert (out / "activities" / "sprint.html").read_text() == "<p>v2</p>"
    assert cfg["activities"]["falta"]["src"] == "activities/falta.html"
    assert not (out / "activities" / "viejo.html").exists()
    html = (out / "index.html").read_text(encoding="utf-8")
    assert cfg["activities"]["sprint"]["version"] in html


def test_config_from_published_roundtrip(tmp_path: Path):
    (tmp_path / "actividades").mkdir()
    (tmp_path / "actividades" / "sprint.html").write_text("<p>v1</p>")
    (tmp_path / "logo.png").write_bytes(b"\x89PNG fake")
    cfg = {
        "theme": "certmind",
        "course": {"title": "Scrum", "logo": "logo.png"},
        "sections": [{"title": "Todo", "from": 1, "to": 3}],
        "links": [{"slide": 1, "rect": {"x": 0, "y": 0, "w": .1, "h": .1}, "href": "https://e.com"}],
        "activities": {"web": "https://e.com/a"},
    }
    out, _ = _publish(tmp_path, cfg)

    rec = config_mod.config_from_published(out)
    assert rec["theme"] == "certmind" and rec["sections"] == cfg["sections"]
    assert rec["course"]["title"] == "Scrum"
    assert Path(rec["course"]["logo"]).is_file()
    assert rec["links"][0]["href"] == "https://e.com" and "kind" not in rec["links"][0]
    assert rec["activities"]["web"]["file"] == "https://e.com/a"
    assert Path(rec["activities"]["sprint"]["file"]).read_text() == "<p>v1</p>"
    # lo recuperado es un config válido que se puede volver a publicar
    config_mod.resolve(rec, rec.get("theme"), 3)


def test_config_from_published_missing(tmp_path: Path):
    assert config_mod.config_from_published(tmp_path / "nada") is None
    (tmp_path / "config.json").write_text("{roto")
    assert config_mod.config_from_published(tmp_path) is None
